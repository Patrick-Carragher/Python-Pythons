from datetime import (
    date,
    datetime,
)
from pathlib import Path

from openpyxl import (
    Workbook,
    load_workbook,
)
from openpyxl.utils.datetime import (
    from_excel,
)

from apps.certification_manager.id_generator import (
    generate_application_id,
    is_valid_application_id,
)
from apps.certification_manager.models import (
    Application,
)


EXCEL_FILE = (
    Path(__file__).parent
    / "data"
    / "applications.xlsx"
)

APPLICATIONS_SHEET = (
    "Applications"
)

HEADERS = [
    "Ankom",
    "Namn",
    "Personnummer",
    "Ny- eller Omcertifiering",
    "Norm",
    "Företag",
    "Adress",
    "Mail",
    "Mail privat",
    "Telefon",
    "Examinationsdatum",
    "Utb hos Säkerhetsbransch?",
    "Plats",
    "Övrigt",
    "Resultat",
    "Status",
    "Application ID",
]


def get_applications_sheet(
    workbook,
):
    if (
        APPLICATIONS_SHEET
        in workbook.sheetnames
    ):
        return workbook[
            APPLICATIONS_SHEET
        ]

    sheet = workbook.active

    sheet.title = (
        APPLICATIONS_SHEET
    )

    return sheet


def format_date_columns(
    sheet,
) -> bool:
    changed = False

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        for column in (
            1,
            11,
        ):
            cell = sheet.cell(
                row=row_number,
                column=column,
            )

            if (
                cell.value is not None
                and
                cell.number_format
                != "yyyy-mm-dd"
            ):
                cell.number_format = (
                    "yyyy-mm-dd"
                )

                changed = True

    return changed


def normalise_excel_date(
    value,
):
    if value is None:
        return None

    if isinstance(
        value,
        datetime,
    ):
        return value.date()

    if isinstance(
        value,
        date,
    ):
        return value

    if isinstance(
        value,
        (int, float),
    ):
        converted = from_excel(
            value
        )

        if isinstance(
            converted,
            datetime,
        ):
            return converted.date()

        if isinstance(
            converted,
            date,
        ):
            return converted

    if isinstance(
        value,
        str,
    ):
        try:
            return date.fromisoformat(
                value
            )

        except ValueError:
            return value

    return value


def get_year_from_value(
    value,
) -> int:
    normalised_date = (
        normalise_excel_date(
            value
        )
    )

    if isinstance(
        normalised_date,
        date,
    ):
        return (
            normalised_date.year
        )

    return (
        datetime.now().year
    )


def ensure_application_ids(
    workbook,
) -> None:
    sheet = get_applications_sheet(
        workbook
    )

    application_id_column = (
        len(HEADERS)
    )

    changed = False

    if (
        sheet.cell(
            row=1,
            column=application_id_column,
        ).value
        != "Application ID"
    ):
        sheet.cell(
            row=1,
            column=application_id_column,
        ).value = (
            "Application ID"
        )

        changed = True

    used_ids: set[str] = set()

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        has_data = any(
            sheet.cell(
                row=row_number,
                column=column,
            ).value
            is not None
            for column in range(
                1,
                application_id_column,
            )
        )

        if not has_data:
            continue

        id_cell = sheet.cell(
            row=row_number,
            column=application_id_column,
        )

        existing_id = (
            str(
                id_cell.value
            ).strip()
            if id_cell.value
            else ""
        )

        if (
            is_valid_application_id(
                existing_id
            )
            and
            existing_id not in used_ids
        ):
            used_ids.add(
                existing_id
            )

            continue

        application_year = (
            get_year_from_value(
                sheet.cell(
                    row=row_number,
                    column=1,
                ).value
            )
        )

        new_application_id = (
            generate_application_id(
                application_year,
                used_ids,
            )
        )

        id_cell.value = (
            new_application_id
        )

        used_ids.add(
            new_application_id
        )

        changed = True

    if format_date_columns(
        sheet
    ):
        changed = True

    if changed:
        workbook.save(
            EXCEL_FILE
        )


def row_to_application(
    row,
) -> Application:
    ankom = (
        normalise_excel_date(
            row[0]
        )
    )

    examinationsdatum = (
        normalise_excel_date(
            row[10]
        )
    )

    return Application(
        ankom=ankom,
        namn=(row[1] or ""),
        personnummer=(row[2] or ""),
        ny_eller_omcertifiering=(row[3] or ""),
        norm=(row[4] or ""),
        foretag=(row[5] or ""),
        adress=(row[6] or ""),
        mail=(row[7] or ""),
        mail_privat=(row[8] or ""),
        telefon=(row[9] or ""),
        examinationsdatum=examinationsdatum,
        utb_hos_sakerhetsbransch=(row[11] or ""),
        plats=(row[12] or ""),
        ovrigt=(row[13] or ""),
        resultat=(row[14] or ""),
        status=(row[15] or ""),
        application_id=(row[16] or ""),
    )


def application_exists(
    application: Application,
) -> bool:
    if (
        not EXCEL_FILE.exists()
        or
        EXCEL_FILE.stat().st_size == 0
    ):
        return False

    workbook = load_workbook(
        EXCEL_FILE
    )

    ensure_application_ids(
        workbook
    )

    sheet = get_applications_sheet(
        workbook
    )

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not any(row):
            continue

        ankom = (
            normalise_excel_date(
                row[0]
            )
        )

        personnummer = row[2]

        certification_type = (
            row[3]
        )

        norm = row[4]

        if (
            personnummer
            == application.personnummer
            and
            certification_type
            == application.ny_eller_omcertifiering
            and
            norm
            == application.norm
            and
            ankom
            == application.ankom
        ):
            return True

    return False


def save_application(
    application: Application,
) -> None:
    if (
        EXCEL_FILE.exists()
        and
        EXCEL_FILE.stat().st_size > 0
    ):
        workbook = load_workbook(
            EXCEL_FILE
        )

        ensure_application_ids(
            workbook
        )

        sheet = get_applications_sheet(
            workbook
        )

    else:
        workbook = Workbook()

        sheet = workbook.active

        sheet.title = (
            APPLICATIONS_SHEET
        )

        sheet.append(
            HEADERS
        )

    existing_ids = {
        str(
            row[16]
        ).strip()
        for row in sheet.iter_rows(
            min_row=2,
            values_only=True,
        )
        if (
            len(row) > 16
            and
            row[16]
        )
    }

    if (
        not is_valid_application_id(
            application.application_id
        )
        or
        application.application_id
        in existing_ids
    ):
        application.application_id = (
            generate_application_id(
                application.ankom.year,
                existing_ids,
            )
        )

    sheet.append(
        [
            application.ankom,
            application.namn,
            application.personnummer,
            application.ny_eller_omcertifiering,
            application.norm,
            application.foretag,
            application.adress,
            application.mail,
            application.mail_privat,
            application.telefon,
            application.examinationsdatum,
            application.utb_hos_sakerhetsbransch,
            application.plats,
            application.ovrigt,
            application.resultat,
            application.status,
            application.application_id,
        ]
    )

    format_date_columns(
        sheet
    )

    workbook.save(
        EXCEL_FILE
    )


def get_all_applications(
) -> list[Application]:
    if (
        not EXCEL_FILE.exists()
        or
        EXCEL_FILE.stat().st_size == 0
    ):
        return []

    workbook = load_workbook(
        EXCEL_FILE
    )

    ensure_application_ids(
        workbook
    )

    sheet = get_applications_sheet(
        workbook
    )

    applications: list[
        Application
    ] = []

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not any(row):
            continue

        applications.append(
            row_to_application(
                row
            )
        )

    return applications


def get_application_by_id(
    application_id: str,
) -> Application | None:
    if (
        not EXCEL_FILE.exists()
        or
        EXCEL_FILE.stat().st_size == 0
    ):
        return None

    workbook = load_workbook(
        EXCEL_FILE
    )

    ensure_application_ids(
        workbook
    )

    sheet = get_applications_sheet(
        workbook
    )

    for row in sheet.iter_rows(
        min_row=2,
        values_only=True,
    ):
        if not any(row):
            continue

        row_application_id = (
            row[16]
            or ""
        )

        if (
            row_application_id
            == application_id
        ):
            return row_to_application(
                row
            )

    return None


def update_application_status_by_id(
    application_id: str,
    status: str,
) -> bool:
    if (
        not EXCEL_FILE.exists()
        or
        EXCEL_FILE.stat().st_size == 0
    ):
        return False

    workbook = load_workbook(
        EXCEL_FILE
    )

    ensure_application_ids(
        workbook
    )

    sheet = get_applications_sheet(
        workbook
    )

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        row_application_id = (
            sheet.cell(
                row=row_number,
                column=17,
            ).value
        )

        if (
            row_application_id
            == application_id
        ):
            sheet.cell(
                row=row_number,
                column=16,
            ).value = status

            workbook.save(
                EXCEL_FILE
            )

            return True

    return False


def delete_application_by_id(
    application_id: str,
) -> bool:
    if (
        not EXCEL_FILE.exists()
        or
        EXCEL_FILE.stat().st_size == 0
    ):
        return False

    workbook = load_workbook(
        EXCEL_FILE
    )

    ensure_application_ids(
        workbook
    )

    sheet = get_applications_sheet(
        workbook
    )

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        row_application_id = (
            sheet.cell(
                row=row_number,
                column=17,
            ).value
        )

        if (
            row_application_id
            == application_id
        ):
            sheet.delete_rows(
                row_number,
                1,
            )

            workbook.save(
                EXCEL_FILE
            )

            return True

    return False
